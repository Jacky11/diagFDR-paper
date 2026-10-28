"""Figure 2B: put the current diagFDR report captures into Figure2.pptx.

Run after `Rscript figure2B_report.R` and `python figure2B_screenshots.py`
(both in this folder; they need diagFDR >= 0.2.1 and ../real_data/out/universes.rds):

    python figure2B_update_pptx.py ../../manuscript/figures/Figure2.pptx out/Figure2.pptx

The three pictures of panel B (title + table of contents; executive summary;
headline table) are replaced by crops of title.png, toc.png, exec.png and
headline.png, keeping panel A (native shapes) untouched. Needs python-pptx and Pillow.
"""
import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

src, dst = sys.argv[1], sys.argv[2]
EMU = 914400

e = Image.open("exec.png")                       # title + alert boxes, without the italic preamble
top, body = e.crop((0, 0, e.width, 100)), e.crop((0, 250, e.width, 1100))
ex = Image.new("RGB", (e.width, top.height + body.height), "white"); ex.paste(top, (0, 0)); ex.paste(body, (0, top.height))
ex.save("exec_crop.png")
h = Image.open("headline.png"); h.crop((0, 0, h.width, 600)).save("headline_crop.png")   # table only
t, toc = Image.open("title.png"), Image.open("toc.png").crop((0, 0, 1100, 1098))
W = 1600; th = int(t.height * W / t.width); t2 = t.resize((W, th), Image.LANCZOS)
tt = Image.new("RGB", (W, th + 20 + toc.height), "white"); tt.paste(t2, (0, 0)); tt.paste(toc, (0, th + 20)); tt.save("title_toc.png")

p = Presentation(src); s = p.slides[0]
def swap(shape, path, width_in=None):
    w, hh = Image.open(path).size
    _, rId = s.part.get_or_add_image_part(path)
    shape._element.xpath(".//a:blip")[0].set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", rId)
    shape.crop_left = shape.crop_right = shape.crop_top = shape.crop_bottom = 0
    if width_in: shape.width = Emu(int(width_in * EMU))
    shape.height = int(shape.width * hh / w)

def walk(sh):
    if sh.shape_type == 6:
        for x in sh.shapes: yield from walk(x)
    else: yield sh
pics = [x for sh in s.shapes for x in walk(sh) if x.shape_type == 13]
pics.sort(key=lambda x: (x.left, x.top))         # left: title/TOC; right: exec (top) then headline
left, right = pics[0], sorted(pics[1:], key=lambda x: x.top)
swap(left, "title_toc.png")
swap(right[0], "exec_crop.png", 4.3); right[0].top = Emu(int(4.05 * EMU))
swap(right[1], "headline_crop.png", 4.3); right[1].left = right[0].left
right[1].top = right[0].top + right[0].height + Emu(int(0.1 * EMU))
assert right[1].top + right[1].height <= p.slide_height, "panel B overflows the slide: reduce the widths"
p.save(dst); print("wrote", dst)
